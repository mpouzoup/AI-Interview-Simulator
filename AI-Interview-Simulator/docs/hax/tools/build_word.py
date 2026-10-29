from pathlib import Path
import json, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

root=Path(__file__).resolve().parents[1]
rows=json.loads((root/'mapping.json').read_text())
refs=json.loads((root/'code-references.json').read_text())
shots=json.loads((root/'screenshots/manifest.json').read_text())
doc=Document()
for border in list(doc.styles.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(.78)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Arial')
 st.element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Arial')
 lang=OxmlElement('w:lang');lang.set(qn('w:val'),'el-GR');st.element.get_or_add_rPr().append(lang)
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.08;normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',25),('Subtitle',13),('Heading 1',16),('Heading 2',13),('Heading 3',11)]:
 doc.styles[name].font.size=Pt(size);doc.styles[name].paragraph_format.space_before=Pt(9);doc.styles[name].paragraph_format.space_after=Pt(6)
 doc.styles[name].paragraph_format.keep_with_next=True
cap=doc.styles['Caption'];cap.font.size=Pt(9.5);cap.font.italic=False;cap.paragraph_format.space_after=Pt(7)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run();r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
doc.core_properties.title='Microsoft HAX Guidelines στην εφαρμογή AI Interview Simulator'
doc.core_properties.subject='Αντιστοίχιση πραγματικού κώδικα και UI με τις 18 οδηγίες'
doc.core_properties.author='AI Interview Simulator';doc.core_properties.keywords='HAX, Human AI Interaction, interview simulator, code review'

def para(text,boldlabel=None,size=None,keep=False):
 p=doc.add_paragraph();p.paragraph_format.keep_together=keep
 if boldlabel:
  p.add_run(boldlabel+' ').bold=True
 p.add_run(text)
 if size:
  for r in p.runs:r.font.size=Pt(size)
 return p

def link(p,label,url):
 rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
 r=OxmlElement('w:r');rp=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'000000');rp.append(c)
 u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(u);r.append(rp);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

# Opening page
h=doc.add_paragraph('Microsoft HAX Guidelines στην εφαρμογή AI Interview Simulator',style='Title')
para('Τεκμηρίωση σχεδιασμού και επιθεώρηση υλοποίησης',size=13)
para('Προς τον επιβλέποντα καθηγητή\n29 Οκτωβρίου 2026',size=10.5)
doc.add_heading('Αντικείμενο και κύριο εύρημα',1)
para('Η εφαρμογή υποστηρίζει τη δομημένη προσομοίωση επτά ερωτήσεων, την παροχή feedback, μία προαιρετική αναθεώρηση και την ανάκαμψη από τεχνικές αποτυχίες. Η αντιστοίχιση με τις 18 οδηγίες Microsoft Human–AI Interaction δείχνει ουσιαστική, αλλά όχι συνολική εφαρμογή. Τα ισχυρότερα παρατηρήσιμα στοιχεία αφορούν τη χρονική οργάνωση της ροής και την εύκολη επίκληση των διαθέσιμων ενεργειών.')
para('Παραμένουν κενά στην εξήγηση των αποφάσεων προς τον χρήστη, στη ρητή περιγραφή της πηγής του feedback, στην αξιολόγηση συγκεκριμένων σχολίων και στην ενημέρωση για μεταβολές δυνατοτήτων. Η αναθεώρηση προσωπικής απάντησης δεν ισοδυναμεί με άμεση διόρθωση AI εξόδου και η αποθήκευση συνομιλίας δεν αποδεικνύει μάθηση.')
doc.add_heading('Οι δύο συνθήκες',1)
para('Σταθερές ελληνικές ερωτήσεις και προκαθορισμένα σχόλια που επιλέγονται από διαφανείς deterministic κανόνες. Δεν χρησιμοποιούνται LLM, embeddings ή παραγωγική AI. Οι HAX αρχές αξιοποιούνται ως πλαίσιο σχεδιασμού της κοινής αλληλεπίδρασης.',boldlabel='Condition A')
para('Οι ίδιες ερωτήσεις και δυναμικό feedback μέσω του υπάρχοντος μηχανισμού Gemini. Στον παρόντα έλεγχο ο εξωτερικός provider αντικαταστάθηκε αποκλειστικά από τοπικό stub. Ελέγχθηκε η πραγματική διεπαφή και η τεχνική ροή, όχι η ποιότητα ζωντανών αποκρίσεων Gemini.',boldlabel='Condition B')
para('Το δεύτερο feedback αφορά μόνο την αναθεωρημένη απάντηση. Δεν συγκρίνει τις δύο εκδοχές και δεν πιστοποιεί βελτίωση.',boldlabel='Κοινή διαδικασία')
para('Το κείμενο είναι τεκμηριωμένη επιθεώρηση της υλοποίησης. Δεν αποτελεί πιστοποίηση της Microsoft, βαθμολογία συμμόρφωσης ή εμπειρική απόδειξη αποτελεσματικότητας.',size=10.5)

# Method + summary page
doc.add_page_break();doc.add_heading('Μέθοδος και συνοπτική αντιστοίχιση',1)
para('Επιθεωρήθηκαν Home.razor, οι feedback engines, ο ελληνικός matcher, οι server μεταβάσεις, το prompt, η επιλογή μοντέλων και ο μηχανισμός recovery. Οι κρίσεις δεν βασίστηκαν στην περιγραφή README. Ελέγχθηκε πραγματικό Blazor UI σε ξεχωριστό Chrome profile, με προσωρινή SQLite και συνθετικές εγγραφές.')
para('Λήφθηκαν 11 screenshots χωρίς αλλαγή του DOM ή του περιεχομένου τους. Το Α ελέγχθηκε έως την έβδομη ερώτηση και με refresh στο στάδιο 3. Στο Β ελέγχθηκε το πρώτο στάδιο με αναθεώρηση, αποτυχία provider και retry. Δεν έγινε εξωτερική κλήση, deployment ή μεταβολή της πραγματικής βάσης. Το αγγλικό SYNTHETIC TEST ONLY είναι fixture, όχι εγκεκριμένη συναίνεση.')
para('Πλήρης: καλύπτεται το σχετικό πεδίο της οδηγίας με άμεσο τεκμήριο. Μερική: υπάρχει συγκεκριμένος μηχανισμός αλλά ουσιαστικός περιορισμός ή μη επαληθευμένη συμπεριφορά. Δεν εφαρμόζεται: δεν εντοπίστηκε σχετική λειτουργία. Δεν είναι σχετικό: η δυνατότητα βρίσκεται εκτός του εγκεκριμένου πεδίου. Οι χαρακτηρισμοί δεν υποκαθιστούν pilot.',size=10.5)
table=doc.add_table(rows=1, cols=4);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
widths=[.62,2.1,2.1,2.1]
for c,w in zip(table.columns,widths):c.width=Inches(w)
for c,t in zip(table.rows[0].cells,['Οδηγία','Συνολικά','Condition A','Condition B']):c.text=t
for r in rows:
 cells=table.add_row().cells
 for cell,value in zip(cells,[f'G{r["number"]}',r['status'],r['conditionA'],r['conditionB']]):cell.text=value
for ri,row in enumerate(table.rows):
 for ci,cell in enumerate(row.cells):
  cell.width=Inches(widths[ci])
  cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  tcPr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'374151' if ri==0 else ('F1F3F5' if ri%2 else 'FFFFFF'));tcPr.append(sh)
  mar=OxmlElement('w:tcMar')
  for edge in ['top','left','bottom','right']:
   e=OxmlElement('w:'+edge);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');mar.append(e)
  tcPr.append(mar)
  for p in cell.paragraphs:
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
   for run in p.runs:run.font.size=Pt(10);run.bold=ri==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
 trpr=table.rows[0]._tr.get_or_add_trPr();trh=OxmlElement('w:tblHeader');trpr.append(trh)
para('Οι διαφορές ανά condition αναλύονται παρακάτω. Δεν υπολογίζεται συνολικό ποσοστό συμμόρφωσης.',size=9.5).paragraph_format.space_before=Pt(6)

# Two detailed assessments per page.
for i,r in enumerate(rows):
 if i%2==0:doc.add_page_break()
 doc.add_heading(f'G{r["number"]} {r["title"]}',1)
 para(r['meaning'],size=10.5)
 para(f'{r["status"]}   |   Α {r["conditionA"]}   |   Β {r["conditionB"]}',boldlabel='Αξιολόγηση',size=10.5)
 para(r['feature'],boldlabel='Στοιχείο εφαρμογής')
 para(r['proof'],boldlabel='Τεκμήριο ελέγχου')
 para(r['limitations'],boldlabel='Περιορισμοί')
 ids=', '.join(s.split('_')[0] for s in r['screenshots']) if r['screenshots'] else 'Κανένα θετικό screenshot λειτουργίας'
 para(ids+'. Κώδικας '+', '.join(r['code'])+'.',boldlabel='Παραπομπές',size=10)
 para(r['nextstep'],boldlabel='Εκκρεμότητα',size=10.5)
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(10);link(p,'Επίσημη πηγή Microsoft G'+str(r['number']),r['url'])
 for run in p.runs:run.font.size=Pt(9)

# Observations and boundaries.
doc.add_page_break();doc.add_heading('Ζητήματα για συζήτηση με τον επιβλέποντα',1)
for label,text in [
 ('Διαφάνεια αυτοματοποίησης','Το UI περιγράφει τι θα κάνει η συνέντευξη αλλά δεν εξηγεί ρητά πώς προκύπτει το feedback ανά condition. Χρειάζεται κοινά συμφωνημένη ενημέρωση που δεν παρουσιάζει το Α ως generative AI και δεν μεταβάλλει ανεξέλεγκτα τις προσδοκίες των ομάδων.'),
 ('Διόρθωση και έλεγχος','Να διατηρηθεί η διάκριση ανάμεσα σε αναθεώρηση απάντησης, retry τεχνικής αποτυχίας, αμφισβήτηση περιεχομένου feedback και ανάκληση πρόσβασης. Το τρέχον σύστημα καλύπτει τα δύο πρώτα και τη λήξη πρόσβασης, όχι άμεση διόρθωση του feedback.'),
 ('Εμπειρική τεκμηρίωση','Χρειάζεται αξιολόγηση κατάλληλων και δύσκολων ελληνικών απαντήσεων για το Α και, μετά από έγκριση πραγματικών κλήσεων, έλεγχος συνάφειας, ύφους και πιθανών διακρίσεων στο Β. Τα screenshots και οι επιτυχείς τεχνικές δοκιμές δεν αποδεικνύουν αυτά τα χαρακτηριστικά.'),
 ('Σταθερότητα της μελέτης','Η επιλογή εξωτερικών μοντέλων και το alias latest δεν παρέχουν από μόνα τους πάγωμα συμπεριφοράς. Απαιτείται καταγραφή εκδόσεων και συμφωνημένη διαχείριση σημαντικών αλλαγών πριν από πραγματική συλλογή.'),
 ('Πεδίο των οδηγιών','Η προσθήκη learning, ratings, εξηγήσεων ή δυνατότητας παράλειψης feedback δεν είναι αυτόματη απαίτηση του HAX. Τέτοιες αλλαγές μπορούν να επηρεάσουν διαδικασία και μετρήσεις και χρειάζονται ξεχωριστή απόφαση.')]:para(text,boldlabel=label)
para('Στον παρόντα κύκλο δεν μεταβλήθηκε η λειτουργικότητα, η κατανομή, οι επτά ερωτήσεις, οι κανόνες Α, το prompt Β ή τα ερευνητικά outcomes. Τα τεκμήρια και τα αποσπάσματα κώδικα παραδίδονται και ως οργανωμένος φάκελος.',size=10.5)

# Screenshots. Original files preserved; only geometric crops in Word.
figure_dir=root/'evidence'/'word-figures';figure_dir.mkdir(exist_ok=True)
caps={
'S01':('Αρχική ενημέρωση και όρια','Κοινό UI Α και Β. Οι οδηγίες περιγράφουν τα επτά στάδια και τα όρια της ανατροφοδότησης. Το fixture συναίνεσης δεν είναι τελικό ερευνητικό κείμενο.','G1 G2 G5 G16'),
'S02':('Ερώτηση και επίκληση της υπηρεσίας','Condition A. Το πεδίο απάντησης και η Αποστολή εμφανίζονται μαζί με το στάδιο. Το ανενεργό κουμπί αντιστοιχεί σε κενή απάντηση.','G3 G7'),
'S03':('Response contingent feedback στο Α','Ερώτηση 3. Το κείμενο περιγράφει πρόβλημα και ενέργεια χωρίς έκβαση. Το προδιατυπωμένο feedback προτείνει διευκρίνιση αποτελέσματος χωρίς βαθμολογία υποψηφίου.','G4 G5 G8 G11 G15 G17'),
'S04':('Μία προαιρετική αναθεώρηση','Condition A. Ο editor έχει το αρχικό κείμενο ως αφετηρία. Η ενέργεια διορθώνει είσοδο του συμμετέχοντα και όχι την ίδια την έξοδο του συστήματος.','G7 G9 G16'),
'S05':('Δεύτερο feedback και συνέχεια','Condition A. Μετά την υποβολή αναθεώρησης παράγεται νέο feedback και προσφέρεται μόνο Συνέχεια. Δεν γίνεται σύγκριση των δύο εκδοχών.','G3 G4 G9 G16'),
'S06':('Επαναφορά μετά από refresh','Condition A. Επανέρχεται ο δεύτερος γύρος της ερώτησης 3. Η έλλειψη παλιότερων σταδίων και πρώτου feedback στην επανασύνθεση περιορίζει τη μνήμη UI.','G12'),
'S07':('Ολοκλήρωση της έβδομης ερώτησης','Condition A. Μετά την έβδομη απάντηση και την επιλογή ολοκλήρωσης εμφανίζεται μήνυμα επιτυχούς τέλους της συνέντευξης. Δεν αποτελεί ολοκλήρωση πραγματικών ερευνητικών ερωτηματολογίων.','G3 G16'),
'S08':('Ουδέτερο fallback σε αβεβαιότητα','Condition A στην ερώτηση 3. Το άσχετο κείμενο δεν οδηγεί σε επινόηση θετικού στοιχείου. Εμφανίζεται η πραγματική ουδέτερη προτροπή οργάνωσης παραδείγματος.','G4 G10'),
'S09':('Διεπαφή feedback στο Β με stub','Condition B. Η διάταξη είναι πραγματική, αλλά το περιεχόμενο προέρχεται από σταθερή τοπική fixture. Δεν τεκμηριώνει την ποιότητα ή συνάφεια πραγματικού Gemini.','G4 G11'),
'S10':('Αποτυχία παραγωγής και επιλογή retry','Condition B με τεχνητή αποτυχία provider. Η αναθεωρημένη απάντηση διατηρείται και ζητείται επανάληψη ανάκτησης. Το marker BROWSER_REVISED αφορά μόνο το test.','G7 G8 G9 G16'),
'S11':('Ανάκαμψη και δεύτερο feedback στο Β','Condition B με stub. Μετά από επανάληψη εμφανίζεται το δεύτερο feedback και μόνο Συνέχεια. Δεν αντικαταστάθηκε η αποτυχία Β από feedback κανόνων Α.','G9 G15')}
limits={'S01':990,'S02':840,'S03':820,'S04':930,'S05':775,'S06':775,'S07':785,'S08':820,'S09':820,'S10':825,'S11':775}
crops=[]
for shot in shots:
 sid=shot['id'].split('_')[0];title,caption,gs=caps[sid]
 doc.add_page_break();doc.add_heading(sid+' '+title,1)
 para(caption,size=10.5)
 original=root/shot['file'];im=Image.open(original)
 box=(85,80,1015,min(limits[sid],im.height));crop=im.crop(box);dest=figure_dir/(sid+'.png');crop.save(dest)
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4)
 image=p.add_run().add_picture(str(dest),width=Inches(6.85))
 image._inline.docPr.set('descr',caption+' Πραγματικό screenshot '+shot['id'])
 para('Σχετικές οδηγίες '+gs+'. Αρχείο '+shot['id']+'.png.',size=9.5)
 para('Γεωμετρική περικοπή πρωτοτύπου για αναγνωσιμότητα. Το πλήρες screenshot και το DOM text περιλαμβάνονται στον φάκελο τεκμηρίωσης. Καμία αντικατάσταση περιεχομένου.',size=9)
 crops.append(dict(id=sid,original=shot['file'],box=box,derived=str(dest.relative_to(root))))
(root/'evidence/word-crops.json').write_text(json.dumps(crops,ensure_ascii=False,indent=2))

# Code references in appendix, independent of machine paths.
doc.add_page_break();doc.add_heading('Παραπομπές στον ελεγμένο κώδικα',1)
para('Οι παραπομπές αφορούν τη μορφή του repository κατά τον έλεγχο της 29ης Οκτωβρίου 2026. Τα αρχεία CODE_REFERENCES.md, code-references.json και evidence/C01 έως C21 περιέχουν πλήρεις διαδρομές, γραμμές, αποσπάσματα και SHA-256. Οι αριθμοί γραμμών μπορεί να μετακινηθούν σε μεταγενέστερη έκδοση.',size=10.5)
for r in refs:
 p=para(r['meaning'],boldlabel=r['id'],size=10.5);p.paragraph_format.keep_with_next=True
 p=para(r['file'].replace('/','/\u200b')+f'   γραμμές {r["start"]} έως {r["end"]}',size=9)
 p.paragraph_format.space_after=Pt(9)

# Primary literature references.
doc.add_page_break();doc.add_heading('Πηγές και ερμηνευτικό πλαίσιο',1)
para('Οι επίσημες αγγλικές ονομασίες προέρχονται από τις σελίδες Microsoft HAX Design Library. Οι ελληνικές αποδόσεις του σκοπού και οι κρίσεις αντιστοίχισης είναι ανάλυση της συγκεκριμένης εφαρμογής, όχι επίσημη ελληνική μετάφραση. Έλεγχος πρόσβασης 29 Οκτωβρίου 2026.')
para('Amershi, S., Weld, D., Vorvoreanu, M., Fourney, A., Nushi, B., Collisson, P., Suh, J., Iqbal, S., Bennett, P. N., Inkpen, K., Teevan, J., Kikin-Gil, R., & Horvitz, E. (2019). Guidelines for Human-AI Interaction. Proceedings of CHI 2019, 1–13. https://doi.org/10.1145/3290605.3300233',size=10.5)
p=doc.add_paragraph();link(p,'Πρωτότυπη δημοσίευση μέσω Microsoft Research','https://www.microsoft.com/en-us/research/wp-content/uploads/2019/01/Guidelines-for-Human-AI-Interaction-camera-ready.pdf')
p=doc.add_paragraph();link(p,'Microsoft HAX Toolkit και οδηγίες','https://www.microsoft.com/en-us/haxtoolkit/ai-guidelines/')
p=doc.add_paragraph();link(p,'Microsoft HAX Design Library','https://www.microsoft.com/en-us/haxtoolkit/library/')
para('Οι οδηγίες κατανέμονται σε τέσσερις χρονικές φάσεις: αρχική αλληλεπίδραση G1–G2, κατά τη χρήση G3–G6, όταν προκύπτει λάθος G7–G11 και σε βάθος χρόνου G12–G18. Το πλαίσιο καθοδηγεί επιλογές σχεδιασμού· η δημοσίευση δεν αξιολογεί την εφαρμογή AI Interview Simulator.')
para('Κάθε επιμέρους αξιολόγηση παραπέμπει στην αντίστοιχη επίσημη ιστοσελίδα. Ο φάκελος SOURCES.md συγκεντρώνει και τις 18 διευθύνσεις. Ούτε η αναφορά οδηγίας ούτε ένα επιτυχημένο screenshot αποτελούν από μόνα τους απόδειξη ακρίβειας, fairness ή μείωσης άγχους.')
doc.save(root/'HAX_AI_Interview_Simulator.docx')
print(root/'HAX_AI_Interview_Simulator.docx')
